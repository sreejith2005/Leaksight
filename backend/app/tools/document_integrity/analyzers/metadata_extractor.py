"""Metadata extraction for Tool B."""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class MetadataExtractor:
    """Extract metadata from supported document formats."""

    _INVOICE_SOFTWARE_KEYWORDS = (
        "adobe acrobat",
        "libreoffice",
        "google docs",
    )

    def extract(self, file_path: Path, doc_type: str) -> dict[str, Any]:
        """Extract metadata without raising on missing fields."""
        metadata: dict[str, Any] = {
            "creation_date": None,
            "modification_date": None,
            "author": None,
            "software": None,
            "page_count": None,
            "revision_count": None,
            "anomalies": [],
        }

        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            metadata.update(self._extract_pdf(file_path))
        elif suffix == ".docx":
            metadata.update(self._extract_docx(file_path))
        elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            metadata.update(self._extract_xlsx(file_path))

        anomalies = list(metadata.get("anomalies", []))
        creation_date = self._parse_iso_date(metadata.get("creation_date"))
        modification_date = self._parse_iso_date(metadata.get("modification_date"))

        if (
            creation_date is not None
            and modification_date is not None
            and modification_date < creation_date
        ):
            anomalies.append("modification_date_precedes_creation_date")

        if modification_date is not None:
            five_years_ago = date.today() - timedelta(days=365 * 5)
            if modification_date < five_years_ago:
                anomalies.append("unusually_old_modification_date")

        author = (metadata.get("author") or "").strip()
        if suffix in {".pdf", ".docx"} and not author:
            anomalies.append("missing_author_metadata")

        software = (metadata.get("software") or "").strip().lower()
        if doc_type == "INVOICE" and any(keyword in software for keyword in self._INVOICE_SOFTWARE_KEYWORDS):
            anomalies.append("software_mismatch_for_invoice")

        metadata["anomalies"] = sorted(set(anomalies))
        if not metadata.get("author"):
            metadata["author"] = None
        if not metadata.get("software"):
            metadata["software"] = None
        return metadata

    def _extract_pdf(self, file_path: Path) -> dict[str, Any]:
        result = {
            "creation_date": None,
            "modification_date": None,
            "author": None,
            "software": None,
            "page_count": None,
            "revision_count": None,
            "anomalies": [],
        }
        try:
            reader = PdfReader(str(file_path))
        except Exception:
            return result

        pdf_metadata = self._safe(lambda: reader.metadata) or {}
        result["creation_date"] = self._normalize_date(
            self._safe(lambda: pdf_metadata.get("/CreationDate"))
        )
        result["modification_date"] = self._normalize_date(
            self._safe(lambda: pdf_metadata.get("/ModDate"))
        )
        result["author"] = self._clean_text(
            self._safe(lambda: pdf_metadata.get("/Author"))
        )
        result["software"] = self._clean_text(
            self._safe(lambda: pdf_metadata.get("/Producer"))
            or self._safe(lambda: pdf_metadata.get("/Creator"))
        )
        result["page_count"] = self._safe(lambda: len(reader.pages))
        return result

    def _extract_docx(self, file_path: Path) -> dict[str, Any]:
        result = {
            "creation_date": None,
            "modification_date": None,
            "author": None,
            "software": None,
            "page_count": None,
            "revision_count": None,
            "anomalies": [],
        }
        try:
            document = DocxDocument(str(file_path))
        except Exception:
            return result

        props = self._safe(lambda: document.core_properties)
        result["creation_date"] = self._normalize_date(
            self._safe(lambda: props.created)
        )
        result["modification_date"] = self._normalize_date(
            self._safe(lambda: props.modified)
        )
        result["author"] = self._clean_text(self._safe(lambda: props.author))
        result["revision_count"] = self._safe_int(self._safe(lambda: props.revision))
        return result

    def _extract_xlsx(self, file_path: Path) -> dict[str, Any]:
        result = {
            "creation_date": None,
            "modification_date": None,
            "author": None,
            "software": None,
            "page_count": None,
            "revision_count": None,
            "anomalies": [],
        }
        try:
            workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        except Exception:
            return result

        props = self._safe(lambda: workbook.properties)
        result["creation_date"] = self._normalize_date(
            self._safe(lambda: props.created)
        )
        result["modification_date"] = self._normalize_date(
            self._safe(lambda: props.modified)
        )
        result["author"] = self._clean_text(
            self._safe(lambda: props.lastModifiedBy)
        )
        return result

    @staticmethod
    def _safe(callback):
        try:
            return callback()
        except Exception:
            return None

    @staticmethod
    def _safe_int(value: Any) -> int | None:
        try:
            if value is None:
                return None
            return int(value)
        except Exception:
            return None

    @staticmethod
    def _clean_text(value: Any) -> str | None:
        try:
            if value is None:
                return None
            text = str(value).strip()
            return text or None
        except Exception:
            return None

    @classmethod
    def _normalize_date(cls, value: Any) -> str | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            if value.tzinfo is None:
                value = value.replace(tzinfo=timezone.utc)
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()

        text = str(value).strip()
        if not text:
            return None

        pdf_match = re.search(r"(\d{4})(\d{2})(\d{2})", text)
        if pdf_match:
            try:
                return date(
                    int(pdf_match.group(1)),
                    int(pdf_match.group(2)),
                    int(pdf_match.group(3)),
                ).isoformat()
            except ValueError:
                return None

        try:
            normalized = text.replace("Z", "+00:00")
            return datetime.fromisoformat(normalized).date().isoformat()
        except Exception:
            return None

    @staticmethod
    def _parse_iso_date(value: Any) -> date | None:
        if not value:
            return None
        try:
            return date.fromisoformat(str(value))
        except Exception:
            return None
